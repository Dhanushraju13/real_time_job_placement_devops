import os
import docx
from docx.shared import Inches, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_TAB_ALIGNMENT
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.oxml import parse_xml
from docx.oxml.ns import nsdecls

def create_evidence_checklist():
    print("Generating Evidence Checklist document...")
    doc = docx.Document()

    doc.core_properties.title = "PartTime - DevOps Laboratory Report Evidence Checklist"
    doc.core_properties.author = "Dhanush R (2303717620521013)"
    doc.core_properties.comments = "Course 23IT723 - DevOps Laboratory, Coimbatore Institute of Technology"

    section = doc.sections[0]
    section.page_width = Inches(8.27)
    section.page_height = Inches(11.69)
    section.top_margin = Inches(0.6)
    section.bottom_margin = Inches(0.6)
    section.left_margin = Inches(0.6)
    section.right_margin = Inches(0.6)

    # Page Border
    pgBorders = parse_xml(f'''
        <w:pgBorders {nsdecls("w")} w:offsetFrom="page">
            <w:top w:val="single" w:sz="4" w:space="24" w:color="auto"/>
            <w:left w:val="single" w:sz="4" w:space="24" w:color="auto"/>
            <w:bottom w:val="single" w:sz="4" w:space="24" w:color="auto"/>
            <w:right w:val="single" w:sz="4" w:space="24" w:color="auto"/>
        </w:pgBorders>
    ''')
    section._sectPr.append(pgBorders)

    # Running Footer
    footer = section.footer
    footer_p = footer.paragraphs[0]
    footer_p.alignment = WD_ALIGN_PARAGRAPH.LEFT
    footer_p.paragraph_format.tab_stops.add_tab_stop(Inches(7.07), WD_TAB_ALIGNMENT.RIGHT)
    
    r_ftr = footer_p.add_run("23IT723 – DevOps Laboratory | Evidence Checklist")
    r_ftr.font.name = "Times New Roman"
    r_ftr.font.size = Pt(8.5)
    r_ftr.font.color.rgb = RGBColor(100, 100, 100)

    footer_p.add_run("\t")
    r_pg = footer_p.add_run()
    r_pg.font.name = "Times New Roman"
    r_pg.font.size = Pt(8.5)
    r_pg.font.color.rgb = RGBColor(80, 80, 80)
    fld1 = parse_xml(f'<w:fldChar {nsdecls("w")} w:fldCharType="begin"/>')
    fld2 = parse_xml(f'<w:instrText {nsdecls("w")} xml:space="preserve"> PAGE </w:instrText>')
    fld3 = parse_xml(f'<w:fldChar {nsdecls("w")} w:fldCharType="separate"/>')
    fld4 = parse_xml(f'<w:fldChar {nsdecls("w")} w:fldCharType="end"/>')
    r_pg._r.append(fld1)
    r_pg._r.append(fld2)
    r_pg._r.append(fld3)
    r_pg._r.append(fld4)

    def set_cell_background(cell, fill_hex):
        tcPr = cell._tc.get_or_add_tcPr()
        shd = parse_xml(f'<w:shd {nsdecls("w")} w:fill="{fill_hex}"/>')
        tcPr.append(shd)

    def set_cell_margins(cell, top=40, bottom=40, left=80, right=80):
        tcPr = cell._tc.get_or_add_tcPr()
        tcMar = parse_xml(f'<w:tcMar {nsdecls("w")}><w:top w:w="{top}" w:type="dxa"/><w:bottom w:w="{bottom}" w:type="dxa"/><w:left w:w="{left}" w:type="dxa"/><w:right w:w="{right}" w:type="dxa"/></w:tcMar>')
        tcPr.append(tcMar)

    def set_table_borders(table, color="B0B0B0", sz="4", val="single"):
        tblPr = table._tbl.tblPr
        borders = parse_xml(f'''
            <w:tblBorders {nsdecls("w")}>
                <w:top w:val="{val}" w:sz="{sz}" w:space="0" w:color="{color}"/>
                <w:bottom w:val="{val}" w:sz="{sz}" w:space="0" w:color="{color}"/>
                <w:left w:val="{val}" w:sz="{sz}" w:space="0" w:color="{color}"/>
                <w:right w:val="{val}" w:sz="{sz}" w:space="0" w:color="{color}"/>
                <w:insideH w:val="{val}" w:sz="{sz}" w:space="0" w:color="{color}"/>
                <w:insideV w:val="{val}" w:sz="{sz}" w:space="0" w:color="{color}"/>
            </w:tblBorders>
        ''')
        tblPr.append(borders)

    def add_p(text="", bold=False, italic=False, size_pt=10, align=WD_ALIGN_PARAGRAPH.JUSTIFY, space_before=0, space_after=3, line_spacing=1.1, color=None):
        p = doc.add_paragraph()
        p.alignment = align
        p.paragraph_format.space_before = Pt(space_before)
        p.paragraph_format.space_after = Pt(space_after)
        p.paragraph_format.line_spacing = line_spacing
        if text:
            r = p.add_run(text)
            r.bold = bold
            r.italic = italic
            r.font.name = "Times New Roman"
            r.font.size = Pt(size_pt)
            if color:
                r.font.color.rgb = color
        return p

    def add_h1(text, space_before=6, space_after=2):
        return add_p(text, bold=True, size_pt=12, align=WD_ALIGN_PARAGRAPH.LEFT, space_before=space_before, space_after=space_after, color=RGBColor(0, 51, 102))

    def add_h2(text, space_before=4, space_after=2):
        return add_p(text, bold=True, size_pt=10.5, align=WD_ALIGN_PARAGRAPH.LEFT, space_before=space_before, space_after=space_after, color=RGBColor(20, 40, 80))

    def add_bullet(text, bold_prefix="", size_pt=9.5, space_after=2):
        p = doc.add_paragraph(style='List Bullet')
        p.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
        p.paragraph_format.space_before = Pt(0)
        p.paragraph_format.space_after = Pt(space_after)
        p.paragraph_format.line_spacing = 1.1
        if bold_prefix:
            r_pre = p.add_run(bold_prefix)
            r_pre.bold = True
            r_pre.font.name = "Times New Roman"
            r_pre.font.size = Pt(size_pt)
        r = p.add_run(text)
        r.font.name = "Times New Roman"
        r.font.size = Pt(size_pt)
        return p

    # Document Header
    add_p("COIMBATORE INSTITUTE OF TECHNOLOGY", bold=True, size_pt=13, align=WD_ALIGN_PARAGRAPH.CENTER, space_before=2, space_after=2, color=RGBColor(0, 51, 102))
    add_p("DEPARTMENT OF INFORMATION TECHNOLOGY | 23IT723 – DEVOPS LABORATORY", bold=True, size_pt=10, align=WD_ALIGN_PARAGRAPH.CENTER, space_before=0, space_after=4)
    add_p("PROJECT EVIDENCE & SUBMISSION VERIFICATION CHECKLIST", bold=True, size_pt=12.5, align=WD_ALIGN_PARAGRAPH.CENTER, space_before=4, space_after=8, color=RGBColor(140, 20, 20))

    # Meta Table
    m_tbl = doc.add_table(rows=2, cols=4)
    m_tbl.alignment = WD_TABLE_ALIGNMENT.CENTER
    set_table_borders(m_tbl)
    widths = [Inches(1.4), Inches(2.13), Inches(1.4), Inches(2.14)]
    
    meta_info = [
        [("Candidate Name:", True), ("Dhanush R", False), ("Register Number:", True), ("2303717620521013", False)],
        [("Semester & Branch:", True), ("VII Semester, B.Tech IT", False), ("Project Title:", True), ("PartTime – Job Portal", False)]
    ]
    for row_idx, row in enumerate(meta_info):
        for col_idx, (text, is_bold) in enumerate(row):
            c = m_tbl.cell(row_idx, col_idx)
            c.width = widths[col_idx]
            if is_bold:
                set_cell_background(c, "F4F6F9")
            set_cell_margins(c, 25, 25, 40, 40)
            p = c.paragraphs[0]
            r = p.add_run(text)
            r.bold = is_bold
            r.font.size = Pt(8.5)

    add_p(
        "Purpose of this Document: In compliance with academic laboratory standards and the strict accuracy rules of the "
        "DevOps Laboratory evaluation, this checklist outlines all required execution artifacts, genuine screenshots, "
        "repository references, and deployment logs that the student must provide and demonstrate prior to the final viva-voce examination."
    , size_pt=9, italic=True, space_before=6, space_after=6)

    # Section 1: Screenshots Checklist
    add_h1("SECTION 1: MANDATORY SCREENSHOT ARTIFACTS CHECKLIST (FIGURES 7.1 TO 7.15)")
    add_p("The report includes 15 clearly defined placeholders. Each placeholder must be replaced with genuine visual evidence:", size_pt=9.5)

    s_tbl = doc.add_table(rows=16, cols=4)
    s_tbl.alignment = WD_TABLE_ALIGNMENT.CENTER
    set_table_borders(s_tbl)
    s_widths = [Inches(0.8), Inches(2.2), Inches(3.07), Inches(1.0)]
    
    s_headers = ["Figure ID", "Figure Caption & Scope", "Exact Capture Command / Source URL", "Student Status"]
    for i, h in enumerate(s_headers):
        c = s_tbl.cell(0, i)
        c.width = s_widths[i]
        set_cell_background(c, "EAECEE")
        set_cell_margins(c, 25, 25, 35, 35)
        p = c.paragraphs[0]
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER if i in [0, 3] else WD_ALIGN_PARAGRAPH.LEFT
        r = p.add_run(h)
        r.bold = True
        r.font.size = Pt(8)

    screenshots_data = [
        ("Figure 7.1", "Home Page & Job Browsing Interface", "Browser at http://localhost/ (Hero banner, search bar, active jobs)", "[  ] Pending"),
        ("Figure 7.2", "Job Search, Filters & Bookmarking", "Browser with 'Weekends' & '₹250-₹400/hr' filters active", "[  ] Pending"),
        ("Figure 7.3", "FastAPI Swagger / OpenAPI Documentation", "Browser at http://localhost:8000/docs or http://localhost/api/docs", "[  ] Pending"),
        ("Figure 7.4", "PostgreSQL Database Tables & Records", "Terminal: 'docker compose exec db psql -U placement -d placement -c \"\\dt\"'", "[  ] Pending"),
        ("Figure 7.5", "GitHub Repository & Commit History", "GitHub repository webpage displaying repo structure & commits", "[  ] Pending"),
        ("Figure 7.6", "PartTime Built Docker Images", "Terminal: 'docker images' showing placement-backend & placement-frontend", "[  ] Pending"),
        ("Figure 7.7", "Running Docker Compose Services", "Terminal: 'docker compose ps' showing all 4 containers healthy/Up", "[  ] Pending"),
        ("Figure 7.8", "Jenkins CI/CD Pipeline Stages", "Jenkins UI: Pipeline Overview or Blue Ocean displaying 9 stages", "[  ] Pending"),
        ("Figure 7.9", "Successful Jenkins Pipeline Execution", "Jenkins UI: Green build status with successful logs", "[  ] Pending"),
        ("Figure 7.10", "Pytest Automated Test Execution", "Terminal: 'python -m pytest backend/tests/ -v' showing tests PASSED", "[  ] Pending"),
        ("Figure 7.11", "Ansible Deployment Playbook Output", "Terminal: 'ansible-playbook -i inventory.ini deploy-parttime.yml'", "[  ] Pending"),
        ("Figure 7.12", "PostgreSQL Docker Volume Persistence", "Terminal: 'docker volume inspect real_time_job_placement_devops_pgdata'", "[  ] Pending"),
        ("Figure 7.13", "Admin Dashboard & Job Moderation", "Browser: Admin dashboard showing analytics & reported jobs queue", "[  ] Pending"),
        ("Figure 7.14", "Webhook-Triggered CI/CD Execution", "GitHub: Repo Settings -> Webhooks delivery log (HTTP 200)", "[  ] Pending"),
        ("Figure 7.15", "Successfully Deployed PartTime Application", "Browser: Full portal accessed via Nginx reverse proxy port 80", "[  ] Pending")
    ]
    for row_idx, data in enumerate(screenshots_data, start=1):
        for col_idx, text in enumerate(data):
            c = s_tbl.cell(row_idx, col_idx)
            c.width = s_widths[col_idx]
            set_cell_margins(c, 18, 18, 30, 30)
            p = c.paragraphs[0]
            p.alignment = WD_ALIGN_PARAGRAPH.CENTER if col_idx in [0, 3] else WD_ALIGN_PARAGRAPH.LEFT
            r = p.add_run(text)
            r.font.size = Pt(8)
            if col_idx == 0:
                r.bold = True

    # Section 2: Technical Evidence Items
    add_h1("SECTION 2: TECHNICAL CONFIGURATION & REPOSITORY EVIDENCE ITEMS")
    tech_checklist = [
        ("Public GitHub Repository URL: ", "Provide actual GitHub repository link and confirm public access for faculty evaluators."),
        ("Git Commit History Evidence: ", "Run 'git log --oneline -n 10' and save output verifying iterative commit history."),
        ("Jenkins Build Console Log: ", "Export raw console text from Jenkins Build #1 (or latest) showing all 9 stages executing."),
        ("Aqua Security Trivy Scan Output: ", "Run 'docker run --rm -v /var/run/docker.sock:/var/run/docker.sock aquasec/trivy:latest image placement-backend:latest' and record output."),
        ("PostgreSQL Schema Dump: ", "Run 'docker compose exec db pg_dump -U placement -s placement' to generate DDL verification snapshot."),
        ("Nginx Gateway Routing Proof: ", "Execute 'curl -I http://localhost/health' and 'curl -I http://localhost/' verifying reverse proxy status headers.")
    ]
    for pfx, bod in tech_checklist:
        add_bullet(bod, bold_prefix=pfx, size_pt=9, space_after=1.5)

    # Section 3: Live Viva-Voce Demonstration Protocol
    add_h1("SECTION 3: LIVE VIVA-VOCE LABORATORY DEMONSTRATION PROTOCOL")
    add_p("To ensure a smooth, confident laboratory examination demonstration, follow this sequence:", size_pt=9.5)
    demo_steps = [
        ("1. Docker Runtime Activation: ", "Launch Docker Desktop. Verify engine status via 'docker info'."),
        ("2. Multi-Container Stack Startup: ", "Execute 'docker compose up -d'. Wait 15 seconds for health checks to turn healthy."),
        ("3. Container Health Verification: ", "Execute 'docker compose ps'. Ensure frontend, backend, db, and nginx show 'Up (healthy)'."),
        ("4. Pytest Test Execution: ", "Open a terminal and run 'docker compose exec -T backend python -m pytest backend/tests/ -v'."),
        ("5. Jenkins Pipeline Demonstration: ", "Open browser to Jenkins (http://localhost:8080). Trigger a build or display past successful pipeline."),
        ("6. Live Portal Demonstration: ", "Navigate to http://localhost/. Demonstrate student registration, search, filters, and 1-click apply.")
    ]
    for pfx, bod in demo_steps:
        add_bullet(bod, bold_prefix=pfx, size_pt=9, space_after=1.5)

    # Signoff Block
    add_h2("Student Verification Sign-Off", space_before=8, space_after=2)
    add_p(
        "I confirm that the above checklist items have been reviewed, and genuine visual and execution artifacts "
        "will be demonstrated to the faculty examiners during the scheduled laboratory viva-voce examination."
    , size_pt=9, italic=True)

    sig_tbl = doc.add_table(rows=1, cols=2)
    sig_tbl.alignment = WD_TABLE_ALIGNMENT.CENTER
    sig_tbl.rows[0].cells[0].width = Inches(3.5)
    sig_tbl.rows[0].cells[1].width = Inches(3.5)
    
    p0 = sig_tbl.cell(0, 0).paragraphs[0]
    p0.add_run("Date: ____________________\nPlace: Coimbatore").font.size = Pt(9)

    p1 = sig_tbl.cell(0, 1).paragraphs[0]
    p1.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    p1.add_run("Student Signature: ___________________________\n").font.size = Pt(9)
    r_name = p1.add_run("DHANUSH R (2303717620521013)")
    r_name.bold = True
    r_name.font.size = Pt(9)

    out_file = "PartTime_DevOps_Lab_Report/PartTime_Report_Evidence_Checklist.docx"
    doc.save(out_file)
    print(f"Successfully saved checklist: {out_file}")
    return out_file

if __name__ == "__main__":
    create_evidence_checklist()
