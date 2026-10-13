import docx
from docx.shared import Inches, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_TAB_ALIGNMENT
from docx.oxml import parse_xml
from docx.oxml.ns import nsdecls

doc = docx.Document()
section = doc.sections[0]
section.page_width = Inches(8.27)
section.page_height = Inches(11.69)
section.top_margin = Inches(0.5)
section.bottom_margin = Inches(0.5)
section.left_margin = Inches(0.5)
section.right_margin = Inches(0.5)

# add page border
pgBorders = parse_xml(f'''
    <w:pgBorders {nsdecls("w")} w:offsetFrom="page">
        <w:top w:val="single" w:sz="4" w:space="24" w:color="auto"/>
        <w:left w:val="single" w:sz="4" w:space="24" w:color="auto"/>
        <w:bottom w:val="single" w:sz="4" w:space="24" w:color="auto"/>
        <w:right w:val="single" w:sz="4" w:space="24" w:color="auto"/>
    </w:pgBorders>
''')
section._sectPr.append(pgBorders)

footer = section.footer
p = footer.paragraphs[0]
p.paragraph_format.tab_stops.add_tab_stop(Inches(7.27), WD_TAB_ALIGNMENT.RIGHT)
r1 = p.add_run("23IT723 – DevOps Laboratory | Final Project")
r1.font.name = "Times New Roman"
r1.font.size = Pt(9)
r2 = p.add_run("\t")
r3 = p.add_run()
fldChar1 = parse_xml(f'<w:fldChar {nsdecls("w")} w:fldCharType="begin"/>')
instrText = parse_xml(f'<w:instrText {nsdecls("w")} xml:space="preserve"> PAGE </w:instrText>')
fldChar2 = parse_xml(f'<w:fldChar {nsdecls("w")} w:fldCharType="separate"/>')
fldChar3 = parse_xml(f'<w:fldChar {nsdecls("w")} w:fldCharType="end"/>')
r3._r.append(fldChar1)
r3._r.append(instrText)
r3._r.append(fldChar2)
r3._r.append(fldChar3)

doc.add_paragraph("Hello world page 1")
doc.add_page_break()
doc.add_paragraph("Hello world page 2")
doc.save("test_out.docx")
print("Saved test_out.docx successfully")
